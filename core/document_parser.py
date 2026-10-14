import io
from typing import BinaryIO


def extract_text_from_upload(uploaded_file) -> str:
    """
    Extracts plain text from a Streamlit UploadedFile or file-like object.
    Supports .pdf, .docx, .txt, .eml.
    """
    if uploaded_file is None:
        return ""

    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if filename.endswith(".pdf"):
        return _extract_from_pdf(io.BytesIO(file_bytes))
    elif filename.endswith(".docx"):
        return _extract_from_docx(io.BytesIO(file_bytes))
    elif filename.endswith(".txt") or filename.endswith(".eml"):
        return _extract_from_text(file_bytes)
    else:
        # Fallback decode
        return _extract_from_text(file_bytes)


def _extract_from_pdf(stream: BinaryIO) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(stream)
        text = ""
        for page in reader.pages:
            page_text = page.extract_text() or ""
            text += page_text + "\n"
        cleaned = text.strip()
        if not cleaned:
            raise ValueError("PDF content appears empty or scanned without OCR.")
        return cleaned
    except Exception as e:
        raise RuntimeError(f"Failed to parse PDF document: {e}")


def _extract_from_docx(stream: BinaryIO) -> str:
    # Try python-docx first
    try:
        import docx
        doc = docx.Document(stream)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if row_cells:
                    paragraphs.append(" | ".join(row_cells))
        return "\n".join(paragraphs).strip()
    except Exception:
        pass

    # Try docx2txt fallback
    try:
        import docx2txt
        text = docx2txt.process(stream)
        if text and text.strip():
            return text.strip()
    except Exception:
        pass

    raise RuntimeError("Could not extract text from DOCX file. Ensure python-docx or docx2txt is installed.")


def _extract_from_text(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            return file_bytes.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore").strip()
